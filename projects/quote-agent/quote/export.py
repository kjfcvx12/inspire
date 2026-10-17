"""견적을 엑셀(.xlsx) 바이트로 만든다. 모든 숫자는 Quote(엔진 결과)에서 그대로 옮긴다. 담당자 검토 전 초안임을 파일에도 적는다."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from quote.engine import Quote, diff

BOLD = Font(bold=True)
WON = "#,##0"


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def to_xlsx(q: Quote, base: Quote) -> bytes:
    r = q.request
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    d = diff(base, q)
    rows = [
        ("초안 — 담당자 검토 전에는 고객에게 보내지 마세요", None),
        ("시나리오", r.label), ("발주", r.client), ("기간(주)", r.weeks),
        ("투입 인원", " · ".join(f"{k} {v}" for k, v in r.team.items())),
        ("단계", ", ".join(r.modules)), ("산출물", ", ".join(r.deliverables)),
        ("신규 항목(LLM 추정)", ", ".join(f"{e.name} {e.effort_pct:g}%" for e in r.extra_items) or "없음"),
        ("공급가액(VAT 별도, 원)", q.supply), ("부가세(원)", q.vat), ("합계(VAT 포함, 원)", q.total),
        ("기준안 대비 차이(VAT 포함, 원)", d["합계(VAT 포함)"]),
        ("필요 대비 투입 공수(%)", round(q.capacity * 100)), ("원가 대비 이윤율(%)", round(q.margin * 100, 1)),
        ("계산 기준", f"가정값 {q.assumptions_version} ({q.assumptions_as_of}) — 공개 정보가 없는 가정 단가"),
    ]
    for row in rows:
        ws.append(row)
    ws["A1"].font = BOLD
    for row in ws.iter_rows(min_row=2):
        row[0].font = BOLD
        if isinstance(row[1].value, int) and "%" not in str(row[0].value) and "주" not in str(row[0].value):
            row[1].number_format = WON
    _widths(ws, (34, 70))

    ws = wb.create_sheet("항목")
    ws.append(("항목", "금액(원)", "기준안 대비(원)"))
    for name, amount in q.lines.items():
        ws.append((name, amount, d.get(name, amount)))
    for c in ws[1]:
        c.font = BOLD
    for row in ws.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = WON
    _widths(ws, (22, 18, 18))

    ws = wb.create_sheet("배분")
    delivs = list(r.deliverables) + (["추정"] if r.extra_items else [])
    ws.append(("단계 / 산출물", *delivs))
    for stage, cells in q.allocation.items():
        ws.append((f"{stage} (LLM 추정)" if stage.startswith("신규:") else stage, *(cells.get(x, 0) for x in delivs)))
    for c in ws[1]:
        c.font = BOLD
    for row in ws.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = WON
    _widths(ws, (34, 18, 18, 18))

    ws = wb.create_sheet("경고")
    ws.append(("경고",))
    ws["A1"].font = BOLD
    for w in q.warnings or ["없음"]:
        ws.append((w,))
    _widths(ws, (100,))

    _force_text(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _force_text(wb: Workbook) -> None:
    """사용자·LLM이 넣은 문자열이 '='·'+'·'-'·'@'로 시작해도 엑셀 수식으로 실행되지 않게 모든 문자열 셀을 텍스트로 고정한다."""
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"
