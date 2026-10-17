from io import BytesIO

from openpyxl import load_workbook

from quote.engine import QuoteRequest, baseline_quote, calculate, sensitivity
from quote.export import to_xlsx


def _row(rows, name):
    return next(r for r in rows if r["변경"] == name)


def test_sensitivity_directions_follow_cost_logic():
    rows = sensitivity(QuoteRequest())
    assert _row(rows, "컨설턴트 +1명")["차이(VAT 포함)"] > 0
    assert _row(rows, "컨설턴트 -1명")["차이(VAT 포함)"] < 0
    assert _row(rows, "기간 +2주")["차이(VAT 포함)"] > 0
    assert _row(rows, "기간 -2주")["차이(VAT 포함)"] < 0


def test_sensitivity_never_removes_last_person_or_week():
    rows = sensitivity(QuoteRequest(team={"PM": 1}, weeks=1))
    names = [r["변경"] for r in rows]
    assert "PM -1명" not in names and "기간 -2주" not in names and "PM +1명" in names


def test_sensitivity_matches_engine():
    req = QuoteRequest()
    r = _row(sensitivity(req), "분석가 +1명")
    expect = calculate(req.model_copy(update={"team": {**req.team, "분석가": req.team["분석가"] + 1}})).total
    assert r["총액(VAT 포함)"] == expect


def test_xlsx_has_expected_sheets_and_matching_total():
    q = calculate(QuoteRequest(client="고려대학교", weeks=12, onsite=True))
    wb = load_workbook(BytesIO(to_xlsx(q, baseline_quote())))
    assert wb.sheetnames == ["요약", "항목", "배분", "경고"]
    flat = [c.value for row in wb["요약"].iter_rows() for c in row if c.value is not None]
    assert q.total in flat and q.supply in flat
    assert any("초안" in str(v) for v in flat)                      # 담당자 검토 전 초안 표시
    assert any(q.assumptions_version in str(v) for v in flat)      # 가정값 버전 기록
    assert {c.value for row in wb["항목"].iter_rows(min_row=2) for c in row[:1]} >= set(q.lines)
